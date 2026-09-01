"""
Bridge report generator.

Reads bridge code_names from an Excel file or a .txt list, fetches info
via bridge_info module, and produces .docx reports by replacing {TAG}
placeholders in a template .docx.

Usage:
    python main.py --bridges bridges.txt --template template.docx
    python main.py --excel bridges.xlsx --template template.docx --out output
"""

from __future__ import annotations

import argparse
import copy
import io
import json
import re
import shutil
from pathlib import Path
from typing import Dict, Iterable

import requests
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Inches, Pt, RGBColor
import docx.opc.constants
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from openpyxl import load_workbook
import pandas as pd

from bridge_info import (
    build_tag_map,
    fetch_bridge,
    fetch_site_ids,
    resensys_login,
)

from sensor_data import process_all_sensor_data
from battery_checker import check_battery
from signal_checker import check_signal
from strain_checker import check_strain
from hpa_checker import check_hpa
from charge_checker import check_charge
from wls_checker import check_wls
from wls_raw_charts import generate_wls_charts
from data_consumption import get_site_data_consumption
import api_bytes_tracker

PLACEHOLDER_RE = re.compile(r"\{([A-Z0-9_]+)\}")

PILOT_REPORTS_DIR = Path("C:/Users/fgkenas/Globiled/Marina Bagourdi - Μεταβλητές/Πιλοτικές Αναφορές")

ELLADA2_LOGO_PATH = Path(__file__).parent / "ellada2" / "Greece 2.0_NextGeneration_gr 1.png"


def add_ellada2_logo(doc, width=Cm(8.9)) -> None:
    """Append the Ellada 2.0 / NextGenerationEU logo as a centered paragraph."""
    if not ELLADA2_LOGO_PATH.exists():
        print(f"  ! Ellada 2.0 logo file not found: {ELLADA2_LOGO_PATH}")
        return
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.add_run().add_picture(str(ELLADA2_LOGO_PATH), width=width)


def _strip_trailing_logo(source) -> None:
    """Remove the trailing Ellada 2.0 logo paragraph (image only, no text) from a doc."""
    body = source.element.body
    for child in reversed(body):
        tag = child.tag.split("}")[-1]
        if tag != "p":
            continue
        texts = "".join(t.text or "" for t in child.iter(qn("w:t"))).strip()
        npic = child.xml.count("<pic:pic")
        if npic and not texts:
            body.remove(child)
        return


def _format_date_el(date_str: str) -> str:
    """Convert YYYY-MM-DD[ HH:MM:SS] to DD/MM/YYYY[ HH:MM:SS]."""
    if not date_str:
        return ""
    try:
        from datetime import datetime
        if " " in date_str and len(date_str) >= 16:
            dt = datetime.strptime(date_str[:19], "%Y-%m-%d %H:%M:%S")
            return dt.strftime("%d/%m/%Y %H:%M:%S")
        dt = datetime.strptime(date_str[:10], "%Y-%m-%d")
        return dt.strftime("%d/%m/%Y")
    except (ValueError, IndexError):
        return date_str


def _sanitize_xml(value: str) -> str:
    """Remove NULL bytes and control characters that break XML parsing."""
    if not value:
        return ""
    import unicodedata
    return "".join(c for c in value if unicodedata.category(c)[0] != "C" or c in "\t\n\r")


# ---------------------------------------------------------------------------
# Cache for bridge data to avoid redundant API calls
# ---------------------------------------------------------------------------
import threading

_cache_lock = threading.Lock()
_bridge_data_cache: dict[str, dict] = {}


# ---------------------------------------------------------------------------
# Utility: parallelized checker batch execution
# ---------------------------------------------------------------------------
def _cached_fetch_bridge(code_name: str) -> dict:
    """Fetch bridge data with thread-safe in-memory cache. Returns {} on failure."""
    with _cache_lock:
        if code_name in _bridge_data_cache:
            return _bridge_data_cache[code_name]
    try:
        data = fetch_bridge(code_name)
    except Exception as exc:
        print(f"     ! fetch_bridge failed: {exc}")
        data = {}
    with _cache_lock:
        _bridge_data_cache[code_name] = data
    return data


def _batch_check_sensors(sensors: list, check_fn, precomputed: dict | None = None) -> dict:
    """Execute check_fn(device_id) in parallel for all sensors.

    If precomputed is provided, use it instead of calling the check function.
    Returns a dict mapping device_id -> check result dict.
    """
    if precomputed is not None:
        return precomputed
    from concurrent.futures import ThreadPoolExecutor, as_completed
    results: dict = {}
    valid_sensors = [(s.get("device_id", ""), s) for s in sensors if s.get("device_id")]
    if not valid_sensors:
        return results
    with ThreadPoolExecutor(max_workers=10) as ex:
        futures = {ex.submit(check_fn, did): did for did, _ in valid_sensors}
        for f in as_completed(futures):
            did = futures[f]
            try:
                results[did] = f.result()
            except Exception:
                results[did] = {"ok": False, "error": "check failed"}
    return results


def _run_all_checkers(sensors: list) -> dict[str, dict]:
    """Run battery, strain, hpa, charge, and wls checkers in parallel.

    Returns dict keyed by checker name, each value is a {did: result} mapping.
    """
    from concurrent.futures import ThreadPoolExecutor, as_completed
    strain_sensors = [s for s in sensors if "strain" in (s.get("type") or "").lower()] if sensors else []
    accel_only = [s for s in sensors if (s.get("type") or "").lower() in ("accelerometer", "accel", "vibration")] if sensors else []
    gw_sensors = [s for s in sensors if (s.get("type") or "").lower() in ("gateway", "gateaway", "gw", "snmx")] if sensors else []
    wl_sensors = [s for s in sensors if (s.get("type") or "").lower() == "waterlevel"] if sensors else []

    results: dict[str, dict] = {}
    with ThreadPoolExecutor(max_workers=5) as ex:
        futures = {
            ex.submit(_batch_check_sensors, sensors, check_battery): "battery",
            ex.submit(_batch_check_sensors, strain_sensors, check_strain): "strain",
            ex.submit(_batch_check_sensors, accel_only, check_hpa): "hpa",
            ex.submit(_batch_check_sensors, gw_sensors, check_charge): "charge",
            ex.submit(_batch_check_sensors, wl_sensors, check_wls): "wls",
        }
        for f in as_completed(futures):
            key = futures[f]
            try:
                results[key] = f.result()
            except Exception:
                results[key] = {}
    return results


def _set_doc_font(doc, font_name: str = "Arial") -> None:
    """Force the document default font so all new runs inherit it."""
    from docx.oxml.ns import qn as _qn
    from docx.oxml import OxmlElement as _OxmlElement

    style_names = ["Normal"] + [f"Heading {i}" for i in range(1, 10)]
    for style_name in style_names:
        try:
            style = doc.styles[style_name]
        except KeyError:
            continue
        rPr = style.element.get_or_add_rPr()
        rFonts = rPr.find(_qn("w:rFonts"))
        if rFonts is None:
            rFonts = _OxmlElement("w:rFonts")
            rPr.insert(0, rFonts)
        # Set direct font names
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rFonts.set(_qn(attr), font_name)
        # Remove theme font attributes that override direct names (e.g. Aptos via majorHAnsi)
        for theme_attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            qn_attr = _qn(theme_attr)
            if rFonts.get(qn_attr) is not None:
                del rFonts.attrib[qn_attr]
        # Force 12pt on heading styles so they match body text
        if style_name.startswith("Heading"):
            sz_val = "24"  # 24 half-points = 12pt
            for sz_tag in ("w:sz", "w:szCs"):
                el = rPr.find(_qn(sz_tag))
                if el is None:
                    el = _OxmlElement(sz_tag)
                    rPr.append(el)
                el.set(_qn("w:val"), sz_val)
        if style_name == "Heading 5":
            style.font.bold = True
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.paragraph_format.space_before = Pt(12)



def shorten_sensor_name(name: str) -> str:
    """Remove first two parts of sensor name (e.g., 'ATT-ATHA-AC-01' -> 'AC-01')."""
    if not name:
        return name
    parts = name.split("-")
    if len(parts) > 2:
        return "-".join(parts[2:])
    return name


class Counter:
    """Simple counter for sequential table and figure numbering."""
    def __init__(self):
        self.table_num = 1  # Start at 1 so first table is 2 (template has Table 1)
        self.figure_num = 0

    def next_table(self) -> int:
        self.table_num += 1
        return self.table_num

    def next_figure(self) -> int:
        self.figure_num += 1
        return self.figure_num





# ---------------------------------------------------------------------------
# Excel input
# ---------------------------------------------------------------------------
def read_code_names(xlsx_path: Path, column: str | int = 1, sheet: str | None = None,
                    skip_header: bool = True) -> list[str]:
    """Read code_name values from a single column of an .xlsx file.

    `column` may be a 1-based int or a column letter (e.g. 'A').
    """
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.active

    if isinstance(column, str) and not column.isdigit():
        col_idx = ws[f"{column}1"].column
    else:
        col_idx = int(column)

    values: list[str] = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if skip_header and i == 0:
            continue
        if col_idx - 1 >= len(row):
            continue
        cell = row[col_idx - 1]
        if cell is None:
            continue
        text = str(cell).strip()
        if text:
            values.append(text)
    return values


# ---------------------------------------------------------------------------
# DOCX placeholder replacement
# ---------------------------------------------------------------------------
def _replace_in_paragraph(paragraph, tags: Dict[str, str]) -> None:
    """Replace {TAG} tokens in a paragraph.

    Merges all run text first (handles placeholders split across runs),
    then writes back to the first run and clears the rest.
    Also applies XML-level replacement to catch any remaining splits.
    """
    # Pass 1: merge runs and replace
    full_text = "".join(run.text for run in paragraph.runs)
    if "{" in full_text:
        def repl(match: re.Match) -> str:
            return tags.get(match.group(1), match.group(0))
        new_text = PLACEHOLDER_RE.sub(repl, full_text)
        if new_text != full_text and paragraph.runs:
            paragraph.runs[0].text = new_text
            for run in paragraph.runs[1:]:
                run.text = ""

    # Pass 2: XML-level replacement for any remaining {TAG} split across XML nodes
    from lxml import etree
    xml_str = etree.tostring(paragraph._element, encoding="unicode")
    if "{" not in xml_str:
        return
    def repl_xml(match: re.Match) -> str:
        return tags.get(match.group(1), match.group(0))
    new_xml = PLACEHOLDER_RE.sub(repl_xml, xml_str)
    if new_xml != xml_str:
        new_el = etree.fromstring(new_xml)
        paragraph._element.getparent().replace(paragraph._element, new_el)


def _iter_paragraphs(doc) -> Iterable:
    for p in doc.paragraphs:
        yield p
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p
    for section in doc.sections:
        for container in (section.header, section.footer):
            for p in container.paragraphs:
                yield p
            for table in container.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            yield p


IMAGE_PLACEHOLDER_RE = re.compile(r"\{BRIDGE_INFO_IMAGE\}")


def _download_image(url: str, timeout: int = 30) -> io.BytesIO | None:
    """Download image from URL, return BytesIO or None on failure."""
    try:
        resp = requests.get(url, timeout=timeout, verify=False)
        resp.raise_for_status()
        buf = io.BytesIO(resp.content)
        buf.seek(0)
        return buf
    except Exception as exc:
        print(f"     ! image download failed ({url}): {exc}")
        return None


def _replace_image_placeholders(doc, tags: Dict[str, str]) -> None:
    """Find {BRIDGE_INFO_IMAGE} placeholders and replace with the actual image.

    If the image cannot be fetched the placeholder is cleared (left empty).
    """
    image_url = tags.get("BRIDGE_INFO_IMAGE", "")
    image_buf = _download_image(image_url) if image_url else None

    for paragraph in _iter_paragraphs(doc):
        full_text = "".join(run.text for run in paragraph.runs)
        if "{BRIDGE_INFO_IMAGE}" not in full_text:
            continue

        # Clear all run text first
        for run in paragraph.runs:
            run.text = ""

        if image_buf:
            image_buf.seek(0)
            run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
            try:
                run.add_picture(image_buf, width=Cm(15))
            except Exception as exc:
                print(f"     ! failed to add image: {exc}")
                run.text = "[Η εικόνα δεν είναι διαθέσιμη]"


def _add_hyperlink(paragraph, url: str, text: str) -> None:
    """Add a hyperlink with display text to a paragraph."""
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)

    run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    # Try Hyperlink style; also add explicit color + underline as fallback
    rStyle = OxmlElement('w:rStyle')
    rStyle.set(qn('w:val'), 'Hyperlink')
    rPr.append(rStyle)
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '0563C1')
    rPr.append(color)
    u = OxmlElement('w:u')
    u.set(qn('w:val'), 'single')
    rPr.append(u)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), '20')  # 10pt * 2
    rPr.append(sz)
    run.append(rPr)

    t = OxmlElement('w:t')
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _set_col_width(cell, width_cm: float) -> None:
    """Set preferred width of a table cell in cm."""
    tcPr = cell._element.get_or_add_tcPr()
    tcW = OxmlElement('w:tcW')
    tcW.set(qn('w:w'), str(int(width_cm * 567)))  # 1 cm = 567 twips
    tcW.set(qn('w:type'), 'dxa')
    # Remove existing tcW if present
    for existing in tcPr.findall(qn('w:tcW')):
        tcPr.remove(existing)
    tcPr.append(tcW)


def add_strain_appendix_section_to_docx(
    doc: Document,
    graphs_dir: Path,
    stats_csv: Path | None = None,
    counter: Counter | None = None,
) -> None:
    """Append Appendix section with remaining strain charts (not in top 25%)."""

    # Add page break before appendix
    p = doc.add_paragraph()
    r = p.add_run()
    r.add_break(docx.enum.text.WD_BREAK.PAGE)

    # Appendix heading
    p = doc.add_paragraph()
    p.style = "Heading 1"
    p.add_run("ΠΑΡΑΡΤΗΜΑ Υπόλοιπες καταγραφές παραμόρφωσης")

    p = doc.add_paragraph(
        "Στο παρόν Παράρτημα παρουσιάζονται οι καταγραφές παραμόρφωσης για τους αισθητήρες "
        "που δεν εμφανίζονται στο κύριο κείμενο (σημείωση: επιλέχθηκαν οι αισθητήρες με τις "
        "μεγαλύτερες απόλυτες τιμές παραμόρφωσης για την κύρια παρουσίαση)."
    )

    # Select remaining strain sensors (not in top 25%)
    remaining_strain_sensors = set()
    if stats_csv and stats_csv.exists():
        df_stats = pd.read_csv(stats_csv, encoding="utf-8-sig")
        df_stats["abs_max"] = df_stats.apply(lambda row: max(abs(row["Ελάχιστη τιμή (με)"]), abs(row["Μέγιστη τιμή (με)"])), axis=1)
        df_sorted = df_stats.sort_values("abs_max", ascending=False)
        n_top = max(1, int(len(df_sorted) * 0.25))
        top_strain_sensors = set(df_sorted.head(n_top)["Αισθητήρας"].str.replace("_strain", "", regex=False))
        all_strain_sensors = set(df_sorted["Αισθητήρας"].str.replace("_strain", "", regex=False))
        remaining_strain_sensors = all_strain_sensors - top_strain_sensors

    if graphs_dir.exists():
        strain_charts = sorted(graphs_dir.glob("*_strain_comp.png"))
        figure_num = counter.next_figure() if counter else 400  # Start figure numbers from 400 for appendix
        for chart_path in strain_charts:
            sensor_name = chart_path.stem.replace("_strain_comp", "")
            # Only show remaining sensors (not in top 25%)
            if remaining_strain_sensors and sensor_name not in remaining_strain_sensors:
                continue
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(chart_path), width=Inches(5.5))
            caption = doc.add_paragraph()
            caption.style = "Caption"
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.paragraph_format.space_before = Pt(0)
            caption.paragraph_format.space_after = Pt(0)
            run = caption.add_run(
                f"Σχήμα {figure_num}: Πρωτογενή δεδομένα παραμόρφωσης (raw) "
                f"και διορθωμένα δεδομένα (temperature-compensated) "
                f"(Αισθητήρας: {sensor_name})"
            )
            run.italic = True
            run.font.size = Pt(10)
            run.font.name = "Arial"
            figure_num += 1

            # Add empty line between plots
            p_space = doc.add_paragraph()
            p_space.paragraph_format.space_after = Pt(12)
            ins(p_space)

            temp_chart = graphs_dir / f"{sensor_name}_temperature.png"
            if temp_chart.exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(temp_chart), width=Inches(5.5))
                caption = doc.add_paragraph()
                caption.style = "Caption"
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = caption.add_run(
                    f"Σχήμα {figure_num}: Καθημερινές μετρήσεις θερμοκρασίας κατά τη διάρκεια "
                    f"περιόδου παρακολούθησης (Αισθητήρας: {sensor_name})"
                )
                run.italic = True
                run.font.size = Pt(10)
                run.font.name = "Arial"
                figure_num += 1

    print(f"  [OK] Added strain appendix section")


def _fix_sensor_name_column(table, col_width_cm: float = 4.5, name_font_pt: int = 9) -> None:
    """Set width of col 0 (sensor name) and reduce its font size to avoid wrapping."""
    for row in table.rows:
        cell = row.cells[0]
        _set_col_width(cell, col_width_cm)
        for para in cell.paragraphs:
            for run in para.runs:
                if run.font.size is None or run.font.size > Pt(name_font_pt):
                    run.font.size = Pt(name_font_pt)
            # If no runs but has text (set via cell.text), re-apply via direct XML
            if not para.runs and para.text:
                from docx.oxml.ns import qn as _qn
                from docx.oxml import OxmlElement as _OxmlElement
                r = _OxmlElement('w:r')
                rPr = _OxmlElement('w:rPr')
                sz = _OxmlElement('w:sz')
                sz.set(_qn('w:val'), str(name_font_pt * 2))
                szCs = _OxmlElement('w:szCs')
                szCs.set(_qn('w:val'), str(name_font_pt * 2))
                r.append(rPr)
                t = _OxmlElement('w:t')
                t.text = para.text
                r.append(t)
                para._p.clear()
                para._p.append(r)


def _fix_all_table_columns(table, font_pt: int = 9) -> None:
    """Reduce font size for all columns in a table to avoid wrapping."""
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    if run.font.size is None or run.font.size > Pt(font_pt):
                        run.font.size = Pt(font_pt)
                # If no runs but has text (set via cell.text), re-apply via direct XML
                if not para.runs and para.text:
                    from docx.oxml.ns import qn as _qn
                    from docx.oxml import OxmlElement as _OxmlElement
                    r = _OxmlElement('w:r')
                    rPr = _OxmlElement('w:rPr')
                    sz = _OxmlElement('w:sz')
                    sz.set(_qn('w:val'), str(font_pt * 2))
                    szCs = _OxmlElement('w:szCs')
                    szCs.set(_qn('w:val'), str(font_pt * 2))
                    r.append(rPr)
                    t = _OxmlElement('w:t')
                    t.text = para.text
                    r.append(t)
                    para._p.clear()
                    para._p.append(r)


def _replace_google_maps_link(doc, tags: Dict[str, str]) -> None:
    """Find {GOOGLE_MAPS_LINK} and replace with hyperlink displaying 'Google maps link'."""
    url = tags.get("GOOGLE_MAPS_LINK", "")
    if not url:
        return

    for paragraph in _iter_paragraphs(doc):
        full_text = "".join(run.text for run in paragraph.runs)
        if "{GOOGLE_MAPS_LINK}" not in full_text:
            continue

        # Clear all runs
        for run in paragraph.runs:
            run.text = ""

        # Add hyperlink with display text
        _add_hyperlink(paragraph, url, "Google maps link")


def render_docx(template_path: Path, tags: Dict[str, str], out_path: Path) -> None:
    sanitized = {k: _sanitize_xml(v) for k, v in tags.items()}
    doc = Document(str(template_path))
    _replace_image_placeholders(doc, sanitized)
    _replace_google_maps_link(doc, sanitized)
    for paragraph in _iter_paragraphs(doc):
        _replace_in_paragraph(paragraph, sanitized)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


# ---------------------------------------------------------------------------
# Orchestration
# ---------------------------------------------------------------------------
def safe_filename(name: str) -> str:
    # Allow ASCII alphanumeric, Greek letters, underscore, hyphen, dot
    safe = re.sub(r"[^a-zA-Z0-9_\-.\u0370-\u03FF]+", "_", name)
    # Replace Greek chars with ASCII equivalents for Windows compatibility
    greek_to_ascii = {
        'Α': 'A', 'Β': 'B', 'Γ': 'G', 'Δ': 'D', 'Ε': 'E', 'Ζ': 'Z', 'Η': 'I', 'Θ': 'TH',
        'Ι': 'I', 'Κ': 'K', 'Λ': 'L', 'Μ': 'M', 'Ν': 'N', 'Ξ': 'X', 'Ο': 'O', 'Π': 'P',
        'Ρ': 'R', 'Σ': 'S', 'Τ': 'T', 'Υ': 'Y', 'Φ': 'F', 'Χ': 'CH', 'Ψ': 'PS', 'Ω': 'O',
        'α': 'a', 'β': 'b', 'γ': 'g', 'δ': 'd', 'ε': 'e', 'ζ': 'z', 'η': 'i', 'θ': 'th',
        'ι': 'i', 'κ': 'k', 'λ': 'l', 'μ': 'm', 'ν': 'n', 'ξ': 'x', 'ο': 'o', 'π': 'p',
        'ρ': 'r', 'σ': 's', 'τ': 't', 'υ': 'y', 'φ': 'f', 'χ': 'ch', 'ψ': 'ps', 'ω': 'o',
    }
    for greek, ascii in greek_to_ascii.items():
        safe = safe.replace(greek, ascii)
    return safe


def generate_bridge_docx(
    code_name: str,
    template: Path,
    out_dir: Path,
    extra_tags: dict | None = None,
    bridge_number: int = 1,
    data: dict | None = None,
    session = None,
    csrftoken: str = "",
    region_code: str = "",
) -> Path | None:
    """Generate a single per-bridge .docx report.

    Fetches data, renders template, appends all sections.
    If data is provided, skip fetch_bridge() and use the pre-fetched data.
    If session/csrftoken are provided, reuse them for WLS charts (avoids re-login).
    If region_code is provided, save to out_dir/region_code/code_name.docx.
    Returns the output path, or None on failure.
    """
    out_dir.mkdir(parents=True, exist_ok=True)
    counter = Counter()

    # Reset per-bridge API bytes tracker to avoid cumulative counts across bridges
    api_bytes_tracker.reset()

    if data is None:
        data = _cached_fetch_bridge(code_name)
    if not data:
        print(f"     ! fetch failed for {code_name}")
        return None

    tags = build_tag_map(data)
    if extra_tags:
        tags.update(extra_tags)

    out_path = out_dir / f"{code_name}.docx"
    try:
        render_docx(template, tags, out_path)
    except Exception as exc:
        print(f"     ! render failed ({type(exc).__name__}): {exc}")
        return None
    print(f"     saved {out_path}")

    safe_name = safe_filename(code_name)
    graphs_dir = Path("graphs") / safe_name
    graphs_dir.mkdir(parents=True, exist_ok=True)

    all_sensors: list = []
    for region in data.get("regions", []):
        for bridge in region.get("bridges", []):
            for site in bridge.get("sites", []):
                for sensor in site.get("sensors", []):
                    sensor["site_id"] = site.get("site_id")
                    all_sensors.append(sensor)

    print(f"     Found {len(all_sensors)} sensors")

    # Fast-skip: bridge has no sensors → save early, no heavy processing
    if not all_sensors:
        print(f"     ! No sensors found for {code_name}, skipping data processing")
        doc = Document(str(out_path))
        _set_doc_font(doc)
        p = doc.add_paragraph("Η γέφυρα δεν διαθέτει αισθητήρες. Δεν υπάρχουν διαθέσιμα δεδομένα.")
        add_ellada2_logo(doc)
        doc.save(str(out_path))
        try:
            dest = PILOT_REPORTS_DIR / region_code if region_code else PILOT_REPORTS_DIR
            dest.mkdir(parents=True, exist_ok=True)
            shutil.copy2(str(out_path), str(dest / out_path.name))
        except Exception as e:
            print(f"     ! Failed to copy to pilot reports dir: {e}")
        return out_path

    # Unified sensor data fetching via REST API (single pass, no duplicates)
    sensor_result: dict = {}
    try:
        sensor_result = process_all_sensor_data(all_sensors, graphs_dir)
    except Exception as exc:
        print(f"     ! sensor data processing failed: {exc}")

    pilot_start = sensor_result.get("pilot_start", "")
    pilot_finish = sensor_result.get("pilot_finish", "")
    # Fallback: extract from sensor metadata if process_all_sensor_data failed
    if not pilot_finish:
        for s in all_sensors:
            pf = s.get("Pilot_Date_Finish", "") or ""
            if pf and not pilot_finish:
                pilot_finish = pf
            ps = s.get("Pilot_Date_Start", "") or ""
            if ps and not pilot_start:
                pilot_start = ps
            if pilot_start and pilot_finish:
                break
    strain_14h_values = sensor_result.get("strain_14h", {})
    hpa_14h_values = sensor_result.get("hpa_14h", {})

    # WLS raw charts (reuse session if provided, otherwise login)
    if all_sensors:
        try:
            wls_session, wls_csrftoken = (session, csrftoken) if session else resensys_login()
            generate_wls_charts(all_sensors, graphs_dir, wls_session, wls_csrftoken)
        except Exception as exc:
            print(f"     ! WLS raw charts failed: {exc}")

    strain_csvs = sorted(graphs_dir.glob("*_strain_stats.csv"))
    accel_csvs = sorted(graphs_dir.glob("*_acceleration_stats.csv"))
    temp_csvs = sorted(graphs_dir.glob("*_temperature_stats.csv"))
    stats_csv = strain_csvs[0] if strain_csvs else graphs_dir / f"{safe_name}_strain_stats.csv"
    accel_stats_csv = accel_csvs[0] if accel_csvs else graphs_dir / f"{safe_name}_acceleration_stats.csv"
    temp_stats_csv = temp_csvs[0] if temp_csvs else graphs_dir / f"{safe_name}_temperature_stats.csv"
    print(f"     stats_csv: {stats_csv} (exists={stats_csv.exists()})")
    print(f"     accel_stats_csv: {accel_stats_csv} (exists={accel_stats_csv.exists()})")
    print(f"     temp_stats_csv: {temp_stats_csv} (exists={temp_stats_csv.exists()})")

    doc = Document(str(out_path))
    _set_doc_font(doc)

    # Pre-fetch all checker results in parallel (avoids sequential API calls)
    checker_results = _run_all_checkers(all_sensors) if all_sensors else {}

    for fn, kwargs in [
        (add_deformation_section_to_docx, dict(doc=doc, graphs_dir=graphs_dir, stats_csv=stats_csv, bridge_name=code_name, counter=counter, bridge_number=bridge_number)),
        (add_temperature_section_to_docx, dict(doc=doc, graphs_dir=graphs_dir, counter=counter, stats_csv=stats_csv, temp_stats_csv=temp_stats_csv, bridge_number=bridge_number)),
        (add_acceleration_section_to_docx, dict(doc=doc, graphs_dir=graphs_dir, accel_stats_csv=accel_stats_csv, bridge_name=code_name, counter=counter, bridge_number=bridge_number)),
        (add_wls_raw_section_to_docx, dict(doc=doc, graphs_dir=graphs_dir, sensors=all_sensors, counter=counter, bridge_number=bridge_number)),
        (add_processed_data_section_to_docx, dict(doc=doc, graphs_dir=graphs_dir, counter=counter, bridge_number=bridge_number)),
        (add_influence_lines_section_to_docx, dict(doc=doc, code_name=code_name, data=data, counter=counter, bridge_number=bridge_number)),
    ]:
        try:
            fn(**kwargs)
        except Exception as exc:
            print(f"     ! {fn.__name__} failed: {exc}")

    bridge_display_name = tags.get("BRIDGE_INFO_NAME", code_name)

    data_usage_mb = ""
    if pilot_start and pilot_finish and all_sensors:
        try:
            unique_site_ids = list(set(s.get("site_id") for s in all_sensors if s.get("site_id")))
            if unique_site_ids:
                total_mb = get_site_data_consumption(pilot_start, pilot_finish, unique_site_ids)
                if total_mb is not None:
                    data_usage_mb = f"{total_mb:.4f}"
                    print(f"     Data consumption (MySQL): {data_usage_mb} MB for sites {unique_site_ids}")
        except Exception as e:
            print(f"     ! MySQL data consumption failed: {e}")
        if not data_usage_mb:
            api_mb = api_bytes_tracker.get_total_mb()
            if api_mb > 0:
                data_usage_mb = f"{api_mb:.1f} (est.)"
                print(f"     Data consumption (API estimate): {data_usage_mb} MB")

    for fn, kwargs in [
        (add_data_devices_section_to_docx, dict(doc=doc, bridge_name=bridge_display_name, data_usage_mb=data_usage_mb, chart_start=pilot_start, chart_end=pilot_finish, sensors=all_sensors, counter=counter, bridge_number=bridge_number)),
        (add_battery_section_to_docx, dict(doc=doc, sensors=all_sensors, counter=counter, bridge_number=bridge_number, precomputed=checker_results.get("battery"), pilot_finish=pilot_finish)),
        # RSSI disabled — server-side bug in /rssi_status endpoint (2025-07-30)
        # (add_signal_section_to_docx, dict(doc=doc, sensors=all_sensors, counter=counter, bridge_number=bridge_number, pilot_finish=pilot_finish)),
        (add_strain_section_to_docx, dict(doc=doc, sensors=all_sensors, counter=counter, bridge_number=bridge_number, strain_14h=strain_14h_values, precomputed=checker_results.get("strain"), pilot_finish=pilot_finish)),
        (add_hpa_section_to_docx, dict(doc=doc, sensors=all_sensors, counter=counter, bridge_number=bridge_number, hpa_14h=hpa_14h_values, precomputed=checker_results.get("hpa"), pilot_finish=pilot_finish)),
        (add_charge_section_to_docx, dict(doc=doc, sensors=all_sensors, counter=counter, bridge_number=bridge_number, precomputed=checker_results.get("charge"), pilot_finish=pilot_finish)),
        (add_wls_section_to_docx, dict(doc=doc, sensors=all_sensors, counter=counter, bridge_number=bridge_number, precomputed=checker_results.get("wls"), pilot_finish=pilot_finish)),
    ]:
        try:
            fn(**kwargs)
        except Exception as exc:
            print(f"     ! {fn.__name__} failed: {exc}")

    add_ellada2_logo(doc)
    doc.save(str(out_path))
    try:
        dest = PILOT_REPORTS_DIR / region_code if region_code else PILOT_REPORTS_DIR
        dest.mkdir(parents=True, exist_ok=True)
        shutil.copy2(str(out_path), str(dest / out_path.name))
    except Exception as e:
        print(f"     ! Failed to copy to pilot reports dir: {e}")
    return out_path


def process(excel: Path, template: Path, out_dir: Path, column: str | int,
            sheet: str | None, skip_header: bool, dump_tags: bool,
            api_code_name: str | None = None) -> None:
    """Read code_names from Excel or API and generate bridge reports."""
    if api_code_name:
        print(f"Fetching site list from API for '{api_code_name}'...")
        try:
            site_ids = fetch_site_ids(api_code_name)
        except Exception as exc:
            print(f"! Failed to fetch site list: {exc}")
            return
        code_names = site_ids
        print(f"Found {len(code_names)} sites from API.")
        # Pre-login for site name lookups (needed when using --api-code-name)
        session, csrftoken = None, ""
        try:
            from bridge_info import get_site_name
            session, csrftoken = resensys_login()
        except Exception as exc:
            print(f"! resensys_login failed (site names may be unavailable): {exc}")
    else:
        code_names = read_code_names(excel, column=column, sheet=sheet,
                                     skip_header=skip_header)
        if not code_names:
            print("No code_names found in Excel.")
            return

    print(f"Found {len(code_names)} bridges.")
    for idx, code_name in enumerate(code_names, start=1):
        print(f"  -> {code_name}")
        bridge_data = None
        if dump_tags:
            try:
                bridge_data = _cached_fetch_bridge(code_name)
                tags = build_tag_map(bridge_data)
                debug_path = out_dir / f"{safe_filename(code_name)}.tags.json"
                debug_path.parent.mkdir(parents=True, exist_ok=True)
                debug_path.write_text(
                    json.dumps(tags, ensure_ascii=False, indent=2),
                    encoding="utf-8",
                )
            except Exception as exc:
                print(f"     ! dump_tags failed: {exc}")
        try:
            generate_bridge_docx(
                code_name=code_name,
                template=template,
                out_dir=out_dir,
                bridge_number=idx,
                data=bridge_data,
            )
        except Exception as exc:
            print(f"     ! generate_bridge_docx failed for {code_name}: {exc}")


def add_deformation_section_to_docx(
    doc: Document,
    graphs_dir: Path,
    stats_csv: Path | None = None,
    bridge_name: str = "",
    counter: Counter | None = None,
    bridge_number: int = 1,
) -> None:
    """Append section 1.1 (deformation) to the document, after the bridge content."""

    # Remove any leftover {DEFORMATION_SECTION} placeholder paragraphs
    body = doc.element.body
    for para in doc.paragraphs:
        full_text = "".join(run.text for run in para.runs) or para.text
        if "{DEFORMATION_SECTION}" in full_text:
            body.remove(para._element)
            break

    # Insert before sectPr (same pattern as temperature section)
    children = list(body)
    anchor_el = children[-2] if children[-1].tag.endswith('}sectPr') and len(children) > 1 else children[-1]

    def ins(new_p):
        nonlocal anchor_el
        anchor_el.addnext(new_p._element)
        anchor_el = new_p._element

    # 1.1 heading
    p_heading = doc.add_paragraph()
    p_heading.style = "Heading 5"
    run = p_heading.add_run(f"{bridge_number}.1 Πρωτογενή δεδομένα (Raw Data)")
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    ins(p_heading)

    # 1.1.1 heading
    p = doc.add_paragraph()
    p.style = "Heading 5"
    run = p.add_run(f"{bridge_number}.1.1 Καταγραφές παραμόρφωσης")
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    ins(p)

    p = doc.add_paragraph(
        "Παρακάτω παρουσιάζονται συνοπτικά οι μετρήσεις παραμορφώσεων που καταγράφηκαν "
        "ανά αισθητήρα κατά την περίοδο της Πιλοτικής Λειτουργίας. Οι καταγραφές συνοψίζονται "
        "στον παρακάτω Πίνακα και στην συνέχεια παρουσιάζονται, ενδεικτικά, οι καταγραφές "
        "που αφορούν στους τέσσερις αισθητήρες με τις υψηλότερες μετρήσεις παραμόρφωσης."
    )
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    ins(p)

    p_cap = doc.add_paragraph()
    p_cap.style = "Caption"
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table_num = counter.next_table() if counter else 32
    p_cap.add_run(f"Πίνακας {table_num}: Σύνοψη μετρήσεων παραμόρφωσης ανά αισθητήρα").italic = True
    ins(p_cap)

    if stats_csv and stats_csv.exists():
        df = pd.read_csv(stats_csv, encoding="utf-8-sig")
        headers = list(df.columns)
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr_cells = table.rows[0].cells
        for i, header in enumerate(headers):
            cell = hdr_cells[i]
            cell.text = header
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.font.bold = True
                    run.font.size = Pt(10)
                    run.font.name = "Arial"
            tcPr = cell._element.get_or_add_tcPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:fill'), 'B4C6E7')
            tcPr.append(shd)
        for _, row in df.iterrows():
            row_cells = table.add_row().cells
            for i, val in enumerate(row):
                cell_text = shorten_sensor_name(str(val)) if i == 0 else str(val)
                row_cells[i].text = cell_text
                for para in row_cells[i].paragraphs:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in para.runs:
                        run.font.size = Pt(10)
                        run.font.name = "Arial"
        anchor_el.addnext(table._tbl)
        anchor_el = table._tbl
    else:
        p = doc.add_paragraph("(Δεν βρέθηκαν δεδομένα στατιστικών)")
        ins(p)

    p = doc.add_paragraph()
    ins(p)

    # Select top 4 strain sensors by max absolute value (or fewer if less than 4 exist)
    top_strain_sensors = set()
    if stats_csv and stats_csv.exists():
        try:
            df_stats = pd.read_csv(stats_csv, encoding="utf-8-sig")
            _min_col = "Ελάχιστη τιμή (με)"
            _max_col = "Μέγιστη τιμή (με)"
            if _min_col in df_stats.columns and _max_col in df_stats.columns:
                df_stats["abs_max"] = df_stats.apply(lambda row: max(abs(row[_min_col]), abs(row[_max_col])), axis=1)
                df_sorted = df_stats.sort_values("abs_max", ascending=False)
                n_top = min(4, len(df_sorted))
                top_strain_sensors = set(df_sorted.head(n_top)["Αισθητήρας"].str.replace("_strain", "", regex=False))
        except Exception as _exc:
            print(f"  ! top sensor selection failed: {_exc}")

    # Check if there are strain charts
    has_charts = False
    if graphs_dir.exists():
        strain_charts = sorted(graphs_dir.glob("*_strain.png"))
        for chart_path in strain_charts:
            sensor_name = chart_path.stem.replace("_strain", "")
            # Only show top 4 sensors in main section
            if top_strain_sensors and sensor_name not in top_strain_sensors:
                continue
            has_charts = True
            break

    if not has_charts:
        p = doc.add_paragraph("Δεν έχει διαγράμματα παραμόρφωσης.")
        ins(p)

    if graphs_dir.exists():
        strain_charts = sorted(graphs_dir.glob("*_strain.png"))
        figure_num = counter.next_figure() if counter else 153
        for chart_path in strain_charts:
            sensor_name = chart_path.stem.replace("_strain", "")
            # Only show top 4 sensors in main section
            if top_strain_sensors and sensor_name not in top_strain_sensors:
                continue
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(chart_path), width=Inches(5.5))
            ins(p)
            caption = doc.add_paragraph()
            caption.style = "Caption"
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.paragraph_format.space_before = Pt(0)
            caption.paragraph_format.space_after = Pt(0)
            run = caption.add_run(
                f"Σχήμα {figure_num}: Πρωτογενή δεδομένα παραμόρφωσης (raw) "
                f"(Αισθητήρας: {sensor_name})"
            )
            run.italic = True
            run.font.size = Pt(10)
            run.font.name = "Arial"
            ins(caption)
            figure_num += 1

            p_space = doc.add_paragraph()
            p_space.paragraph_format.space_after = Pt(12)
            ins(p_space)

    # Add page break after section
    add_page_break(doc)

    print(f"  [OK] Added deformation section")


def add_temperature_section_to_docx(
    doc: Document,
    graphs_dir: Path,
    counter: Counter | None = None,
    stats_csv: Path | None = None,
    temp_stats_csv: Path | None = None,
    bridge_number: int = 1,
) -> None:
    """Append 1.1.2 temperature section after the last element in the document body."""
    body = doc.element.body
    children = list(body)
    anchor_el = children[-2] if children[-1].tag.endswith('}sectPr') and len(children) > 1 else children[-1]

    def ins(new_p):
        nonlocal anchor_el
        anchor_el.addnext(new_p._element)
        anchor_el = new_p._element

    p = doc.add_paragraph()
    p.style = "Heading 5"
    run = p.add_run(f"{bridge_number}.1.2 Καταγραφές θερμοκρασίας")
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    ins(p)

    p = doc.add_paragraph(
        "Παρακάτω παρουσιάζονται οι καθημερινές μετρήσεις θερμοκρασίας που καταγράφηκαν κατά την περίοδο της Πιλοτικής Λειτουργίας. "
        "Οι καταγραφές συνοψίζονται στον παρακάτω Πίνακα. Οι τιμές αφορούν τις καταγραφές από το εσωτερικό θερμόμετρο κάθε αισθητήρα παραμόρφωσης. "
        "Οι καταγραφές που δείχνονται αφορούν τους τέσσερις αισθητήρες με τις υψηλότερες μετρήσεις παραμόρφωσης."
    )
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    ins(p)

    p = doc.add_paragraph()
    ins(p)

    # Add temperature statistics table for all strain sensors
    if temp_stats_csv and temp_stats_csv.exists():
        df_temp_stats = pd.read_csv(temp_stats_csv, encoding="utf-8-sig")
        headers = list(df_temp_stats.columns)

        p_cap = doc.add_paragraph()
        p_cap.style = "Caption"
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        table_num = counter.next_table() if counter else 2
        p_cap.add_run(f"Πίνακας {table_num}: Σύνοψη μετρήσεων θερμοκρασίας για κάθε αισθητήρα παραμόρφωσης").italic = True
        ins(p_cap)

        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr_cells = table.rows[0].cells
        for i, header in enumerate(headers):
            cell = hdr_cells[i]
            cell.text = header
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.font.bold = True
                    run.font.size = Pt(10)
                    run.font.name = "Arial"
            tcPr = cell._element.get_or_add_tcPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:fill'), 'B4C6E7')
            tcPr.append(shd)
        for _, row in df_temp_stats.iterrows():
            row_cells = table.add_row().cells
            for i, val in enumerate(row):
                cell_text = shorten_sensor_name(str(val)) if i == 0 else str(val)
                row_cells[i].text = cell_text
                for para in row_cells[i].paragraphs:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in para.runs:
                        run.font.size = Pt(10)
                        run.font.name = "Arial"
        anchor_el.addnext(table._tbl)
        anchor_el = table._tbl

    p = doc.add_paragraph()
    ins(p)

    # Build top-sensor filter matching section 1.1.1 (top 4 sensors)
    top_strain_sensors: set[str] = set()
    if stats_csv and stats_csv.exists():
        try:
            df_stats = pd.read_csv(stats_csv, encoding="utf-8-sig")
            _min_col = "Ελάχιστη τιμή (με)"
            _max_col = "Μέγιστη τιμή (με)"
            if _min_col in df_stats.columns and _max_col in df_stats.columns:
                df_stats["abs_max"] = df_stats.apply(
                    lambda row: max(abs(row[_min_col]), abs(row[_max_col])), axis=1
                )
                df_sorted = df_stats.sort_values("abs_max", ascending=False)
                n_top = min(4, len(df_sorted))
                top_strain_sensors = set(df_sorted.head(n_top)["Αισθητήρας"].str.replace("_strain", "", regex=False))
        except Exception as _exc:
            print(f"  ! top sensor selection failed: {_exc}")

    # Check if there are temperature charts
    has_charts = False
    if graphs_dir.exists():
        temp_charts = sorted(graphs_dir.glob("*_temperature.png"))
        for chart_path in temp_charts:
            sensor_name = chart_path.stem.replace("_temperature", "")
            if top_strain_sensors and sensor_name not in top_strain_sensors:
                continue
            has_charts = True
            break

    if not has_charts:
        p = doc.add_paragraph("Δεν έχει διαγράμματα θερμοκρασίας.")
        ins(p)

    if graphs_dir.exists():
        temp_charts = sorted(graphs_dir.glob("*_temperature.png"))
        figure_num = counter.next_figure() if counter else 160
        for chart_path in temp_charts:
            sensor_name = chart_path.stem.replace("_temperature", "")
            if top_strain_sensors and sensor_name not in top_strain_sensors:
                continue
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(chart_path), width=Inches(5.5))
            ins(p)
            caption = doc.add_paragraph()
            caption.style = "Caption"
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = caption.add_run(
                f"Σχήμα {figure_num}: Καθημερινές μετρήσεις θερμοκρασίας κατά τη διάρκεια "
                f"περιόδου παρακολούθησης (Αισθητήρας: {sensor_name})"
            )
            run.italic = True
            run.font.size = Pt(10)
            run.font.name = "Arial"
            ins(caption)
            figure_num += 1

            # Add empty line between plots
            p_space = doc.add_paragraph()
            p_space.paragraph_format.space_after = Pt(12)
            ins(p_space)

    # Add page break after section
    add_page_break(doc)

    print(f"  [OK] Added temperature section")


def add_acceleration_section_to_docx(
    doc: Document,
    graphs_dir: Path,
    accel_stats_csv: Path | None = None,
    bridge_name: str = "",
    counter: Counter | None = None,
    bridge_number: int = 1,
) -> None:
    """Append 1.1.3 acceleration section after the last element in the document body."""
    body = doc.element.body
    children = list(body)
    # anchor = last non-sectPr element
    anchor_el = children[-2] if children[-1].tag.endswith('}sectPr') and len(children) > 1 else children[-1]

    def ins(new_p):
        nonlocal anchor_el
        anchor_el.addnext(new_p._element)
        anchor_el = new_p._element

    p = doc.add_paragraph()
    p.style = "Heading 5"
    run = p.add_run(f"{bridge_number}.1.3 Καταγραφές επιτάχυνσης")
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    ins(p)

    p = doc.add_paragraph(
        "Στον παρακάτω πίνακα παρουσιάζεται η μέγιστη παρατηρούμενη επιτάχυνση κατά τον άξονα Ζ, καθώς και ο αντίστοιχος μέσος όρος και η ημερήσια διακύμανση."
    )
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    ins(p)

    p_cap = doc.add_paragraph()
    p_cap.style = "Caption"
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table_num = counter.next_table() if counter else 33
    p_cap.add_run(f"Πίνακας {table_num}: Μέγιστη επιτάχυνση  κατά τον άξονα Ζ και συνολική μέγιστη επιτάχυνση ανά συσκευή").italic = True
    ins(p_cap)

    if accel_stats_csv and accel_stats_csv.exists():
        df = pd.read_csv(accel_stats_csv, encoding="utf-8-sig")
        headers = list(df.columns)
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr_cells = table.rows[0].cells
        for i, header in enumerate(headers):
            cell = hdr_cells[i]
            cell.text = header
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.font.bold = True
                    run.font.size = Pt(10)
                    run.font.name = "Arial"
            tcPr = cell._element.get_or_add_tcPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:fill'), 'B4C6E7')
            tcPr.append(shd)
        for _, row in df.iterrows():
            row_cells = table.add_row().cells
            for i, val in enumerate(row):
                cell_text = shorten_sensor_name(str(val)) if i == 0 else str(val)
                row_cells[i].text = cell_text
                for para in row_cells[i].paragraphs:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        anchor_el.addnext(table._tbl)
        anchor_el = table._tbl
    else:
        p = doc.add_paragraph("(Δεν βρέθηκαν δεδομένα επιτάχυνσης)")
        ins(p)

    p = doc.add_paragraph()
    ins(p)

    p = doc.add_paragraph(
        "Οι μη επεξεργασμένες καταγραφές της επιτάχυνσης δίνονται παρακάτω για κάθε συσκευή για τους τρείς άξονες. Οι επιταχυνσιογράφοι, βρίσκονται σε κατάσταση αναμονής (idle mode) και ενεργοποιούνται κάθε 5 ώρες, εκτός υπάρχει υπέρβαση του προκαθορισμένου επιπέδου ενεργοποίησης (trigger level)."
    )
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    ins(p)

    if graphs_dir.exists():
        accel_charts = sorted([p for p in graphs_dir.glob("*_acceleration.png") if not p.name.startswith('_')])
        figure_num = counter.next_figure() if counter else 163
        for chart_path in accel_charts:
            stem = chart_path.stem.replace("_acceleration", "")
            # stem may be e.g. "SENSOR_NAME_hpax" or "SENSOR_NAME_hpay"
            for _tag in ("_hpax", "_hpay", "_hpaz"):
                if stem.endswith(_tag):
                    axis_label = _tag.lstrip("_").upper()  # HPAX / HPAY / HPAZ
                    axis_label = axis_label[:3] + "-" + axis_label[3:]  # HPA-X / HPA-Y / HPA-Z
                    sensor_name = stem[: -len(_tag)] + f" ({axis_label})"
                    break
            else:
                continue  # skip charts without HPA-X/Y/Z axis tag (remnants)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(chart_path), width=Inches(5.5))
            ins(p)
            caption = doc.add_paragraph()
            caption.style = "Caption"
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = caption.add_run(
                f"Σχήμα {figure_num}: Σήμα συμβάντος μέγιστης επιτάχυνσης (Αισθητήρας: {sensor_name})"
            )
            run.italic = True
            run.font.size = Pt(10)
            run.font.name = "Arial"
            ins(caption)
            figure_num += 1

            # Add empty line between plots
            p_space = doc.add_paragraph()
            p_space.paragraph_format.space_after = Pt(12)
            ins(p_space)

    # Add page break after section
    add_page_break(doc)

    print(f"  [OK] Added acceleration section")


def add_wls_raw_section_to_docx(
    doc: Document,
    graphs_dir: Path,
    sensors: list | None = None,
    counter: Counter | None = None,
    bridge_number: int = 1,
) -> None:
    """Append section 1.1.4 (water level raw data charts) after the last element in the document body.
    Shows charts for waterlevel sensors; if none exist, inserts an empty-section notice.
    """
    body = doc.element.body
    children = list(body)
    anchor_el = children[-2] if children[-1].tag.endswith('}sectPr') and len(children) > 1 else children[-1]

    def ins(new_p):
        nonlocal anchor_el
        anchor_el.addnext(new_p._element)
        anchor_el = new_p._element

    p = doc.add_paragraph()
    p.style = "Heading 5"
    run = p.add_run(f"{bridge_number}.1.4 Καταγραφές στάθμης ύδατος")
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    ins(p)

    wls_sensors = [
        s for s in (sensors or [])
        if (s.get("type") or "").lower() == "waterlevel"
    ]

    if not wls_sensors:
        p = doc.add_paragraph("Η γέφυρα δεν έχει αισθητήρες μέτρησης της στάθμης υδάτων.")
        ins(p)
        print(f"  [OK] Added WLS raw section (no sensors)")
        return

    p = doc.add_paragraph(
        "Παρακάτω παρουσιάζονται οι καταγραφές στάθμης υδάτων που καταγράφηκαν "
        "κατά την περίοδο της Πιλοτικής Λειτουργίας."
    )
    ins(p)

    p = doc.add_paragraph()
    ins(p)

    if graphs_dir.exists():
        wls_charts = sorted(graphs_dir.glob("*_wls.png"))
        figure_num = counter.next_figure() if counter else 1
        for chart_path in wls_charts:
            sensor_name = chart_path.stem.replace("_wls", "")
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(chart_path), width=Inches(5.5))
            ins(p)
            caption = doc.add_paragraph()
            caption.style = "Caption"
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = caption.add_run(
                f"Σχήμα {figure_num}: Καταγραφή στάθμης υδάτων (Αισθητήρας: {sensor_name})"
            )
            run.italic = True
            run.font.size = Pt(10)
            run.font.name = "Arial"
            ins(caption)
            figure_num += 1

            # Add empty line between plots
            p_space = doc.add_paragraph()
            p_space.paragraph_format.space_after = Pt(12)
            ins(p_space)

    # Add page break after section
    add_page_break(doc)

    print(f"  [OK] Added WLS raw section")


def add_processed_data_section_to_docx(
    doc: Document,
    graphs_dir: Path,
    counter: Counter | None = None,
    bridge_number: int = 1,
) -> None:
    """Append section 1.2 (processed/algorithmic results) after the last element in the document body."""
    body = doc.element.body
    children = list(body)
    anchor_el = children[-2] if children[-1].tag.endswith('}sectPr') and len(children) > 1 else children[-1]

    def ins(new_p):
        nonlocal anchor_el
        anchor_el.addnext(new_p._element)
        anchor_el = new_p._element

    # Section 1.2
    p = doc.add_paragraph()
    p.style = "Heading 4"
    run = p.add_run(f"{bridge_number}.2 Αποτελέσματα επεξεργασμένων μετρήσεων με χρήση αλγορίθμων")
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    ins(p)

    p = doc.add_paragraph(
        "Κατά την διάρκεια της περιόδου παρακολούθησης αναγνωρίζονται τα «Γεγονότα» (events), που έλαβαν "
        "χώρα ως τμήματα των σημάτων επιτάχυνσης κάθε αισθητήρα. Τα κομμάτια αυτά του σήματος αναγνωρίζονται "
        "εφαρμόζοντας έναν αλγόριθμο που εντοπίζει τα ενεργά τμήματα με ένα κινούμενο παράθυρο (moving window) "
        "που αναγνωρίζει τα σημεία που στο σήμα υπάρχει σημαντική ενέργεια."
    )
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    ins(p)

    # Section 1.2.1
    p = doc.add_paragraph()
    p.style = "Heading 5"
    run = p.add_run(f"{bridge_number}.2.1 Καταγραφές παραμόρφωσης")
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    ins(p)

    # Select top 4 strain sensors by max absolute value (or fewer if less than 4 exist)
    top_strain_sensors = set()
    # Read stats from CSV to determine top sensors
    stats_csv_path = None
    if graphs_dir.exists():
        strain_csvs = sorted(graphs_dir.glob("*_strain_stats.csv"))
        stats_csv_path = strain_csvs[0] if strain_csvs else None
    
    if stats_csv_path and stats_csv_path.exists():
        try:
            df_stats = pd.read_csv(stats_csv_path, encoding="utf-8-sig")
            _min_col = "Ελάχιστη τιμή (με)"
            _max_col = "Μέγιστη τιμή (με)"
            if _min_col in df_stats.columns and _max_col in df_stats.columns:
                df_stats["abs_max"] = df_stats.apply(
                    lambda row: max(abs(row[_min_col]), abs(row[_max_col])), axis=1
                )
                df_sorted = df_stats.sort_values("abs_max", ascending=False)
                n_top = min(4, len(df_sorted))
                top_strain_sensors = set(df_sorted.head(n_top)["Αισθητήρας"].str.replace("_strain", "", regex=False))
        except Exception as _exc:
            print(f"  ! top sensor selection failed: {_exc}")

    # Check if there are strain comp charts
    has_charts = False
    if graphs_dir.exists():
        comp_charts = sorted(graphs_dir.glob("*_strain_comp.png"))
        for chart_path in comp_charts:
            sensor_name = chart_path.stem.replace("_strain_comp", "")
            # Only show top 4 sensors
            if top_strain_sensors and sensor_name not in top_strain_sensors:
                continue
            has_charts = True
            break

    if has_charts:
        p = doc.add_paragraph(
            "Στα παρακάτω διαγράμματα παρουσιάζονται οι χρονοσειρές των μετρήσεων παραμόρφωσης και η "
            "αντίστοιχη διόρθωσή τους ως προς την θερμοκρασία. Από τα διαγράμματα έχει αφαιρεθεί η τιμή της "
            "παραμόρφωσης αναφοράς ε0, η οποία λήφθηκε ίση με την τιμή της παραμόρφωσης λίγο πριν την "
            "έναρξη της δοκιμαστικής φόρτισης. Για την διόρθωση ως προς την θερμοκρασία θεωρήθηκε ως θερμοκρασία αναφοράς, η τιμή Τ0, που ήταν η θερμοκρασία κατά την δοκιμαστική φόρτιση. "
            "Τα διορθωμένα σήματα έχουν ομαλοποιηθεί με την χρήση του κυλιόμενου μέσου (rolling average). Η σύγκριση των δύο καμπυλών επιτρέπει την αξιολόγηση της "
            "ποιότητας των αρχικών δεδομένων και αναδεικνύει την επίδραση της θερμοκρασίας και του θορύβου στις καταγραφές των αισθητήρων"
        )
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        ins(p)
    else:
        p = doc.add_paragraph("Δεν έχει διαγράμματα παραμόρφωσης.")
        ins(p)

    # Add strain comp charts (only top 4 sensors)
    if graphs_dir.exists():
        comp_charts = sorted(graphs_dir.glob("*_strain_comp.png"))
        figure_num = counter.next_figure() if counter else 173
        for chart_path in comp_charts:
            sensor_name = chart_path.stem.replace("_strain_comp", "")
            # Only show top 4 sensors
            if top_strain_sensors and sensor_name not in top_strain_sensors:
                continue
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(chart_path), width=Inches(5.5))
            ins(p)
            caption = doc.add_paragraph()
            caption.style = "Caption"
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = caption.add_run(
                f"Σχήμα {figure_num}: Καταγραφή παραμόρφωσης με και χωρίς διόρθωση θερμοκρασίας (Αισθητήρας: {sensor_name})"
            )
            run.italic = True
            run.font.size = Pt(10)
            run.font.name = "Arial"
            ins(caption)
            figure_num += 1

            # Add empty line between plots
            p_space = doc.add_paragraph()
            p_space.paragraph_format.space_after = Pt(12)
            ins(p_space)

    # Section 1.2.2
    p = doc.add_paragraph()
    p.style = "Heading 5"
    run = p.add_run(f"{bridge_number}.2.2 Καταγραφές επιτάχυνσης")
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    ins(p)

    # Check if there are acceleration charts
    has_charts = False
    if graphs_dir.exists():
        highrate_charts = sorted(graphs_dir.glob("*_accel_highrate.png"))
        for chart_path in highrate_charts:
            stem = chart_path.stem
            if "_hpaz" in stem.lower():
                has_charts = True
                break

    if has_charts:
        p = doc.add_paragraph(
            "Κατά την διάρκεια της περιόδου παρακολούθησης αναγνωρίζονται τα «Γεγονότα» (events), που έλαβαν "
            "χώρα ως τμήματα των σημάτων επιτάχυνσης που είχαμε υπέρβαση ενός κατωφλιού που ορίζεται ως το επίπεδο "
            "ενεργοποίησης (trigger level) κάθε αισθητήρα. Τα κομμάτια αυτά του σήματος αναγνωρίζονται "
            "εφαρμόζοντας έναν αλγόριθμο που εντοπίζει τα ενεργά τμήματα με ένα κινούμενο παράθυρο (moving window) "
            "που αναγνωρίζει τα σημεία που στο σήμα υπάρχει σημαντική ενέργεια. "
            "Στα παρακάτω διαγράμματα παρουσιάζονται οι καταγραφές επιτάχυνσης υψηλής συχνότητας ανά αισθητήρα κατά την περίοδο της Πιλοτικής Λειτουργίας. "
            "Για κάθε αισθητήρα σημειώνεται η κορυφαία τιμή επιτάχυνσης (Peak) που καταγράφηκε. "
            "Εξετάζονται μόνο οι καταγραφές κατά τον άξονα Ζ που αντιστοιχούν σε φορτία κυκλοφορίας και δείχνονται τα τέσσερα σημαντικότερα συμβάντα Γεγονότα, για κάθε αισθητήρα επιτάχυνσης."
        )
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        ins(p)
    else:
        p = doc.add_paragraph("Δεν έχει διαγράμματα επιτάχυνσης.")
        ins(p)

    if graphs_dir.exists():
        highrate_charts = sorted(graphs_dir.glob("*_accel_highrate.png"))
        figure_num_accel = counter.next_figure() if counter else 227
        
        # Group charts by sensor and limit to max 4 events per sensor
        sensor_events = {}
        for chart_path in highrate_charts:
            stem = chart_path.stem
            if "_hpaz" not in stem.lower():
                continue
            sensor_name = re.sub(r"_accel_highrate$", "", stem)
            display_name = re.sub(r"_hpa[xyz].*", "", sensor_name, flags=re.IGNORECASE)
            evt_match = re.search(r"_evt(\d+)", sensor_name, re.IGNORECASE)
            evt_num = int(evt_match.group(1)) if evt_match else 0
            
            if display_name not in sensor_events:
                sensor_events[display_name] = []
            sensor_events[display_name].append((evt_num, chart_path, sensor_name))
        
        # Sort events by event number and limit to 4 per sensor
        for display_name, events in sensor_events.items():
            events.sort(key=lambda x: x[0])
            events = events[:4]  # Limit to max 4 events
            
            for evt_num, chart_path, sensor_name in events:
                evt_label = f" (Γεγονός {evt_num + 1})"
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(chart_path), width=Inches(5.5))
                ins(p)
                caption = doc.add_paragraph()
                caption.style = "Caption"
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = caption.add_run(
                    f"Σχήμα {figure_num_accel}: Καταγραφή επιτάχυνσης υψηλής συχνότητας (Αισθητήρας: {display_name}{evt_label})"
                )
                run.italic = True
                run.font.size = Pt(10)
                run.font.name = "Arial"
                ins(caption)
                figure_num_accel += 1

    # Add page break after section
    add_page_break(doc)

    print(f"  [OK] Added processed data section")


def _insert_after_paragraph(anchor_para, new_element) -> None:
    """Insert an XML element immediately after anchor_para in the document body."""
    anchor_para._element.addnext(new_element)


def add_data_devices_section_to_docx(
    doc: Document,
    bridge_name: str = "",
    data_usage_mb: str = "",
    chart_start: str = "",
    chart_end: str = "",
    sensors: list | None = None,
    counter: Counter | None = None,
    bridge_number: int = 1,
) -> None:
    """Append section 1.4 (data consumption + sensor table) at end of document body."""

    # Count device types from bridge API sensor metadata
    strain_count = 0
    accel_count = 0
    gw_count = 0
    
    if sensors:
        # Count devices by type from the bridge API sensor metadata (no extra API calls)
        for s in sensors:
            stype = (s.get("type") or "").lower()
            code_name = (s.get("Code_name") or s.get("code_name") or "").lower()
            if stype == "strain" or "strain" in stype:
                strain_count += 1
            elif stype in ("accelerometer", "accel") or "-ac-" in code_name:
                accel_count += 1
            elif stype in ("gw", "snmx", "snmx gateway", "gateway", "gateaway"):
                gw_count += 1

    print(f"     DEBUG sensors section: strain={strain_count} accel={accel_count} gw={gw_count} total_sensors={len(sensors) if sensors else 0}")

    mb_text = data_usage_mb if data_usage_mb else "______"
    start_text = _format_date_el(chart_start) if chart_start else "______"
    end_text = _format_date_el(chart_end) if chart_end else "______"
    bridge_display = bridge_name if bridge_name else ""

    p = doc.add_paragraph()
    p.style = "Heading 5"
    run = p.add_run(f"{bridge_number}.4 Κατανάλωση όγκου δεδομένων και κατάσταση συσκευών")
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)

    p = doc.add_paragraph()
    p.style = "Heading 5"
    run = p.add_run(f"{bridge_number}.4.1 Κατανάλωση όγκου δεδομένων")
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)

    p = doc.add_paragraph(
        f"Στη {bridge_display} χρησιμοποιήθηκαν {mb_text} MB κυψελοειδών δεδομένων κατά τη διάρκεια "
        f"της περιόδου παρακολούθησης μεταξύ {start_text} και {end_text} UTC."
    )
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for run in p.runs:
        run.font.size = Pt(11)
        run.font.name = "Arial"

    p_cap = doc.add_paragraph()
    p_cap.style = "Caption"
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table_num = counter.next_table() if counter else 36
    p_cap.add_run(f"Πίνακας {table_num}: Εγκατεστημένοι αισθητήρες ανά τύπο").italic = True

    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(["Τύπος αισθητήρων", "Αριθμός αισθητήρων"]):
        cell = hdr_cells[i]
        cell.text = header
        for para in cell.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                run.font.bold = True
                run.font.size = Pt(10)
        tcPr = cell._element.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), 'B4C6E7')
        tcPr.append(shd)
    for label, count in [
        ("SenSpot™ Wireless 1D/3D Vibration", str(accel_count)),
        ("SenSpot™ Wireless Strain Gauge (Event Detection)", str(strain_count)),
        ("SeniMax™ Cellular Gateway", str(gw_count)),
    ]:
        row_cells = table.add_row().cells
        row_cells[0].text = label
        row_cells[1].text = count
        for cell in row_cells:
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Ensure table is before sectPr, add trailing paragraph so Word renders table
    body = doc.element.body
    sect_pr = body.find(qn('w:sectPr'))
    if sect_pr is not None:
        sect_pr.addprevious(table._tbl)
    doc.add_paragraph()

    # Add page break after section
    add_page_break(doc)

    print(f"  [OK] Added data/devices section")


def add_influence_lines_section_to_docx(
    doc: Document,
    code_name: str = "",
    data: Dict[str, Any] | None = None,
    counter: Counter | None = None,
    bridge_number: int = 1,
) -> None:
    """Append section i.3 / 1.3 (Γραμμές επιρροής) at end of document body.

    For each influence line returned by the bridge API:
      - PlotCaption is inserted as normal text
      - Plot is downloaded and inserted as a centred image
      - ReportText is used as the caption, with the placeholder "Σχήμα Χ"
        replaced by the sequential figure number within this section.
    """

    p = doc.add_paragraph()
    p.style = "Heading 4"
    run = p.add_run(f"{bridge_number}.3 Γραμμές επιρροής")
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)

    if not data:
        data = fetch_bridge(code_name) if code_name else {}

    influence_lines: Dict[str, Any] = {}
    for region in data.get("regions", []):
        for bridge in region.get("bridges", []):
            il = bridge.get("influence_lines")
            if isinstance(il, dict) and il:
                influence_lines = il
                break
        if influence_lines:
            break

    if not influence_lines:
        p = doc.add_paragraph("Λόγω ιδιαιτερότητας του στατικού φορέα, κρίθηκε σκόπιμο να μην παρουσιαστούν Γραμμές Επιρροής για αυτή την γέφυρα.")
        print(f"  [OK] Added influence lines section (no data)")
        return

    def _clean_text(text: str) -> str:
        if not text:
            return ""
        text = text.replace("\r\n", " ").replace("\n", " ")
        text = text.strip().lstrip("\ufeff")
        return re.sub(r"\s+", " ", text).strip()

    figure_num = 1
    for line_key in sorted(influence_lines.keys()):
        line = influence_lines[line_key]
        if not isinstance(line, dict):
            continue

        plot_caption = _clean_text(line.get("PlotCaption", ""))
        plot_url = line.get("Plot", "")
        report_text = _clean_text(line.get("ReportText", ""))

        if plot_caption:
            p = doc.add_paragraph(plot_caption)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(6)
            for run in p.runs:
                run.font.size = Pt(10)
                run.font.name = "Arial"

        image_buf = _download_image(plot_url) if plot_url else None
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if image_buf:
            image_buf.seek(0)
            p.add_run().add_picture(image_buf, width=Inches(5.5))
        else:
            p.add_run("[Η εικόνα δεν είναι διαθέσιμη]")

        if report_text:
            caption_text = re.sub(
                r"Σχήμα\s+Χ",
                f"Σχήμα {figure_num}",
                report_text,
                count=1,
            )
            caption = doc.add_paragraph()
            caption.style = "Caption"
            caption.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            caption.paragraph_format.space_before = Pt(0)
            caption.paragraph_format.space_after = Pt(6)
            run = caption.add_run(caption_text)
            run.italic = True
            run.font.size = Pt(10)
            run.font.name = "Arial"
            doc.add_paragraph()

        doc.add_paragraph()

        figure_num += 1

    # Add page break after section
    add_page_break(doc)

    print(f"  [OK] Added influence lines section")


def add_battery_section_to_docx(
    doc: Document,
    sensors: list | None = None,
    counter: Counter | None = None,
    bridge_number: int = 1,
    precomputed: dict | None = None,
    pilot_finish: str = "",
) -> None:
    """Append section 1.4.2 (battery status table) at end of document body."""
    from datetime import datetime

    p = doc.add_paragraph()
    p.style = "Heading 5"
    run = p.add_run(f"{bridge_number}.4.2 Σύνοψη κατάστασης μπαταρίας")
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)

    p_cap = doc.add_paragraph()
    p_cap.style = "Caption"
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table_num = counter.next_table() if counter else 50
    p_cap.add_run(f"Πίνακας {table_num}: Κατάσταση και τελευταίες μετρήσεις μπαταρίας").italic = True

    # Create table with header row
    headers = ["Αισθητήρας", "Site", "Ώρα", "Μπαταρία (V)", "Status", "Σχόλια"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header row
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        cell = hdr_cells[i]
        cell.text = header
        for para in cell.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                run.font.bold = True
                run.font.size = Pt(10)
        tcPr = cell._element.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), 'B4C6E7')
        tcPr.append(shd)

    # Add row for each sensor with battery data
    if sensors:
        print(f"     DEBUG battery section: processing {len(sensors)} sensors")
        battery_results = _batch_check_sensors(sensors, check_battery, precomputed=precomputed)
        for sensor in sensors:
            device_id = sensor.get("device_id", "")
            # Try multiple fields for sensor name from bridge API
            sensor_name = (sensor.get("Code_name", "") or 
                          sensor.get("name", "") or 
                          sensor.get("sensor_name", "") or
                          sensor.get("sensor_code", "") or
                          sensor.get("device_id", ""))
            site_id = sensor.get("site_id", "")
            print(f"     DEBUG sensor: device_id={device_id}, name={sensor_name}, site_id={site_id}")

            if device_id:
                result = battery_results.get(device_id, {})
                if result.get("ok") and result.get("metrics", {}).get("Battery Voltage"):
                    voltage = result.get("metrics", {}).get("Battery Voltage", "")
                    status = "ΚΑΛΗ"
                    _explain = result.get("explain", "")
                    _upper_words = re.findall(r'\b[Α-ΩΆΈΉΊΌΎΏΪΫ]{2,}\b', _explain)
                    comments = " ".join(_upper_words) if _upper_words else _explain
                    latest_time = result.get("latest_time", "")
                elif result.get("ok"):
                    voltage = "—"
                    status = "ΚΑΛΗ"
                    comments = "Η μπαταρία είναι εντός ορίων (χωρίς ακριβή τιμή)"
                    latest_time = result.get("latest_time", "")
                else:
                    voltage = "—"
                    status = "N/A"
                    comments = result.get("error", "Σφάλμα ελέγχου")
                    latest_time = ""
            else:
                voltage = ""
                status = "N/A"
                comments = "Δεν υπάρχει device_id"
                latest_time = ""

            # Use pilot finish date at 14:00
            if pilot_finish:
                latest_time = _format_date_el(pilot_finish) + " 14:00"
            else:
                latest_time = _format_date_el(latest_time) or datetime.now().strftime("%d/%m/%Y 14:00")

            # Fallback to device_id if sensor_name is still empty
            if not sensor_name:
                sensor_name = device_id

            row_cells = table.add_row().cells
            row_cells[0].text = shorten_sensor_name(sensor_name)
            row_cells[1].text = site_id
            row_cells[2].text = latest_time
            # Ensure voltage is rounded to 3 decimals (in case checker didn't apply rounding)
            if voltage and isinstance(voltage, (int, float)):
                voltage = f"{float(voltage):.3f}"
            row_cells[3].text = str(voltage) if voltage else ""
            row_cells[4].text = status
            row_cells[5].text = comments

            for cell in row_cells:
                for para in cell.paragraphs:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Ensure table is before sectPr, add trailing paragraph
    body = doc.element.body
    sect_pr = body.find(qn('w:sectPr'))
    if sect_pr is not None:
        sect_pr.addprevious(table._tbl)
    doc.add_paragraph()

    _fix_all_table_columns(table)
    
    # Add page break after section
    add_page_break(doc)
    
    print(f"  [OK] Added battery section")


def add_signal_section_to_docx(
    doc: Document,
    sensors: list | None = None,
    counter: Counter | None = None,
    bridge_number: int = 1,
    pilot_finish: str = "",
) -> None:
    """Append section 1.4.3 (RSSI signal table) at end of document body."""
    from datetime import datetime

    p = doc.add_paragraph()
    p.style = "Heading 5"
    run = p.add_run(f"{bridge_number}.4.3 Ποιότητα σήματος")
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)

    p_cap = doc.add_paragraph()
    p_cap.style = "Caption"
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table_num = counter.next_table() if counter else 51
    p_cap.add_run(f"Πίνακας {table_num}: Τελευταία μέτρηση RSSI ανά συσκευή").italic = True

    # Create table with header row
    headers = ["Αισθητήρας", "Site", "Ώρα", "RSSI (dBm)", "Status", "Notes"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header row
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        cell = hdr_cells[i]
        cell.text = header
        for para in cell.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                run.font.bold = True
                run.font.size = Pt(10)
        tcPr = cell._element.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), 'B4C6E7')
        tcPr.append(shd)

    # Add row for each sensor with RSSI data
    if sensors:
        print(f"     DEBUG signal section: processing {len(sensors)} sensors")
        signal_results = _batch_check_sensors(sensors, check_signal)
        for sensor in sensors:
            device_id = sensor.get("device_id", "")
            # Try multiple fields for sensor name from bridge API
            sensor_name = (sensor.get("Code_name", "") or 
                           sensor.get("name", "") or 
                           sensor.get("sensor_name", "") or
                           sensor.get("sensor_code", "") or
                           sensor.get("device_id", ""))
            site_id = sensor.get("site_id", "")
            print(f"     DEBUG sensor: device_id={device_id}, name={sensor_name}, site_id={site_id}")

            if device_id:
                result = signal_results.get(device_id, {})
                if result.get("ok"):
                    rssi = result.get("metrics", {}).get("RSSI", "")
                    status = "Καλό"
                    _explain = result.get("explain", "")
                    _upper_words = re.findall(r'\b[Α-ΩΆΈΉΊΌΎΏΪΫ]{2,}\b', _explain)
                    comments = " ".join(_upper_words) if _upper_words else _explain
                    latest_time = result.get("latest_time", "")
                else:
                    rssi = ""
                    status = "N/A"
                    comments = result.get("error", "Σφάλμα ελέγχου")
                    latest_time = ""
            else:
                rssi = ""
                status = "N/A"
                comments = "Δεν υπάρχει device_id"
                latest_time = ""

            # Use pilot finish date if no specific timestamp available
            if not latest_time:
                latest_time = pilot_finish or datetime.now().strftime("%d/%m/%Y")

            # Fallback to device_id if sensor_name is still empty
            if not sensor_name:
                sensor_name = device_id

            row_cells = table.add_row().cells
            row_cells[0].text = shorten_sensor_name(sensor_name)
            row_cells[1].text = site_id
            row_cells[2].text = latest_time
            # Ensure rssi is rounded to 3 decimals
            if rssi and isinstance(rssi, (int, float)):
                rssi = f"{float(rssi):.3f}"
            row_cells[3].text = str(rssi) if rssi else ""
            row_cells[4].text = status
            row_cells[5].text = comments

            for cell in row_cells:
                for para in cell.paragraphs:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Ensure table is before sectPr, add trailing paragraph
    body = doc.element.body
    sect_pr = body.find(qn('w:sectPr'))
    if sect_pr is not None:
        sect_pr.addprevious(table._tbl)
    doc.add_paragraph()

    _fix_all_table_columns(table)
    
    # Add page break after section
    add_page_break(doc)
    
    print(f"  [OK] Added signal section")


def add_strain_section_to_docx(
    doc: Document,
    sensors: list | None = None,
    counter: Counter | None = None,
    bridge_number: int = 1,
    strain_14h: dict | None = None,
    precomputed: dict | None = None,
    pilot_finish: str = "",
) -> None:
    """Append section 1.4.4 (strain quality table) at end of document body.
    Only processes sensors with type 'strain' (SG sensors).

    If strain_14h is provided (from temp-compensated chart generation),
    uses the 14:00 pilot-end-day values instead of calling check_strain API.
    """
    from datetime import datetime


    p = doc.add_paragraph()
    p.style = "Heading 5"
    run = p.add_run(f"{bridge_number}.4.4 Ποιότητα μετρήσεων παραμόρφωσης")
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)

    p_cap = doc.add_paragraph()
    p_cap.style = "Caption"
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table_num = counter.next_table() if counter else 52
    p_cap.add_run(f"Πίνακας {table_num}: Τελευταίες τιμές παραμόρφωσης και δείκτες ποιότητας").italic = True

    # Create table with header row
    headers = ["Αισθητήρας", "Site", "Time", "Τελευταία μέτρηση (με)", "Status", "Σχόλια"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header row
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        cell = hdr_cells[i]
        cell.text = header
        for para in cell.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                run.font.bold = True
                run.font.size = Pt(10)
        tcPr = cell._element.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), 'B4C6E7')
        tcPr.append(shd)

    # Filter for strain sensors (SG sensors) and add row for each
    strain_sensors = [s for s in sensors if "strain" in (s.get("type") or "").lower()] if sensors else []
    
    if strain_sensors:
        print(f"     DEBUG strain section: processing {len(strain_sensors)} strain sensors")
        use_14h = strain_14h is not None and len(strain_14h) > 0
        if not use_14h:
            strain_results = _batch_check_sensors(strain_sensors, check_strain, precomputed=precomputed)
        for sensor in strain_sensors:
            device_id = sensor.get("device_id", "")
            # Try multiple fields for sensor name from bridge API
            sensor_name = (sensor.get("Code_name", "") or 
                           sensor.get("name", "") or 
                           sensor.get("sensor_name", "") or
                           sensor.get("sensor_code", "") or
                           sensor.get("device_id", ""))
            site_id = sensor.get("site_id", "")

            if device_id:
                if use_14h and sensor_name in strain_14h:
                    v14 = strain_14h[sensor_name]
                    strain_value = v14.get("strain", "")
                    latest_time = v14.get("timestamp", "")
                    status = "Καλή"
                    comments = f"Μέτρηση 14:00 (compensated)"
                else:
                    result = strain_results.get(device_id, {}) if not use_14h else {}
                    if result.get("ok"):
                        strain_value = result.get("metrics", {}).get("Strain", "")
                        status = "Καλή"
                        _explain = result.get("explain", "")
                        _upper_words = re.findall(r'\b[Α-ΩΆΈΉΊΌΎΏΪΫ]{2,}\b', _explain)
                        comments = " ".join(_upper_words) if _upper_words else _explain
                        latest_time = result.get("latest_time", "")
                    else:
                        strain_value = ""
                        status = "N/A"
                        comments = result.get("error", "Σφάλμα ελέγχου")
                        latest_time = ""
            else:
                strain_value = ""
                status = "N/A"
                comments = "Δεν υπάρχει device_id"
                latest_time = ""

            # Use pilot finish date at 14:00
            if pilot_finish:
                latest_time = _format_date_el(pilot_finish) + " 14:00"
            else:
                latest_time = _format_date_el(latest_time) or datetime.now().strftime("%d/%m/%Y 14:00")

            # Fallback to device_id if sensor_name is still empty
            if not sensor_name:
                sensor_name = device_id

            row_cells = table.add_row().cells
            row_cells[0].text = shorten_sensor_name(sensor_name)
            row_cells[1].text = site_id
            row_cells[2].text = latest_time
            # Ensure strain values are rounded to 3 decimals (handle both string and numeric)
            try:
                if strain_value:
                    strain_value = f"{float(strain_value):.3f}"
            except (ValueError, TypeError):
                pass
            row_cells[3].text = str(strain_value) if strain_value else ""
            row_cells[4].text = status
            row_cells[5].text = comments

            for cell in row_cells:
                for para in cell.paragraphs:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Ensure table is before sectPr, add trailing paragraph
    body = doc.element.body
    sect_pr = body.find(qn('w:sectPr'))
    if sect_pr is not None:
        sect_pr.addprevious(table._tbl)
    doc.add_paragraph()

    _fix_all_table_columns(table)
    
    # Add page break after section
    add_page_break(doc)
    
    print(f"  [OK] Added strain section")


def add_hpa_section_to_docx(
    doc: Document,
    sensors: list | None = None,
    counter: Counter | None = None,
    bridge_number: int = 1,
    hpa_14h: dict | None = None,
    precomputed: dict | None = None,
    pilot_finish: str = "",
) -> None:
    """Append section 1.4.5 (HPA/accelerometer quality table) at end of document body.
    Only processes sensors with type 'Accelerometer' (AC sensors).

    If hpa_14h is provided (from accel highrate chart generation),
    uses the 14:00 pilot-end-day values instead of calling check_hpa API.
    """
    from datetime import datetime


    p = doc.add_paragraph()
    p.style = "Heading 5"
    run = p.add_run(f"{bridge_number}.4.5 Ποιότητα μετρήσεων επιτάχυνσης")
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)

    p_cap = doc.add_paragraph()
    p_cap.style = "Caption"
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table_num = counter.next_table() if counter else 53
    p_cap.add_run(f"Πίνακας {table_num}: Τελευταίες μετρήσεις και κατάσταση επιταχυνσιογράφου [mg(1/1000g)]").italic = True

    # Create table with header row
    headers = ["Αισθητήρας", "Site", "Ώρα", "Τελευταία μέτρηση", "Status", "Σχόλια"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header row
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        cell = hdr_cells[i]
        cell.text = header
        for para in cell.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                run.font.bold = True
                run.font.size = Pt(10)
        tcPr = cell._element.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), 'B4C6E7')
        tcPr.append(shd)

    # Filter for accelerometer sensors (AC sensors) and add row for each
    accel_sensors = [s for s in sensors if (s.get("type") or "").lower() in ("accelerometer", "accel", "vibration")] if sensors else []
    
    if accel_sensors:
        print(f"     DEBUG HPA section: processing {len(accel_sensors)} accelerometer sensors")
        use_14h = hpa_14h is not None and len(hpa_14h) > 0
        if not use_14h:
            hpa_results = _batch_check_sensors(accel_sensors, check_hpa, precomputed=precomputed)
        for sensor in accel_sensors:
            device_id = sensor.get("device_id", "")
            # Try multiple fields for sensor name from bridge API
            sensor_name = (sensor.get("Code_name", "") or 
                           sensor.get("name", "") or 
                           sensor.get("sensor_name", "") or
                           sensor.get("sensor_code", "") or
                           sensor.get("device_id", ""))
            site_id = sensor.get("site_id", "")

            if device_id:
                if use_14h and sensor_name in hpa_14h:
                    v14 = hpa_14h[sensor_name]
                    hpa_value = v14.get("hpa", "")
                    latest_time = v14.get("timestamp", "")
                    status = "Καλή"
                    comments = f"Μέτρηση 14:00 (raw)"
                else:
                    result = hpa_results.get(device_id, {}) if not use_14h else {}
                    if result.get("ok"):
                        hpa_value = result.get("metrics", {}).get("HPA", "")
                        status = "Καλή"
                        _explain = result.get("explain", "")
                        _upper_words = re.findall(r'\b[Α-ΩΆΈΉΊΌΎΏΪΫ]{2,}\b', _explain)
                        comments = " ".join(_upper_words) if _upper_words else _explain
                        latest_time = result.get("latest_time", "")
                    else:
                        hpa_value = ""
                        status = "N/A"
                        comments = result.get("error", "Σφάλμα ελέγχου")
                        latest_time = ""
            else:
                hpa_value = ""
                status = "N/A"
                comments = "Δεν υπάρχει device_id"
                latest_time = ""

            # Use pilot finish date at 14:00
            if pilot_finish:
                latest_time = _format_date_el(pilot_finish) + " 14:00"
            else:
                latest_time = _format_date_el(latest_time) or datetime.now().strftime("%d/%m/%Y 14:00")

            # Fallback to device_id if sensor_name is still empty
            if not sensor_name:
                sensor_name = device_id

            row_cells = table.add_row().cells
            row_cells[0].text = shorten_sensor_name(sensor_name)
            row_cells[1].text = site_id
            row_cells[2].text = latest_time
            # Ensure hpa is rounded to 3 decimals
            if hpa_value and isinstance(hpa_value, (int, float)):
                hpa_value = f"{float(hpa_value):.3f}"
            row_cells[3].text = str(hpa_value) if hpa_value else ""
            row_cells[4].text = status
            row_cells[5].text = comments

            for cell in row_cells:
                for para in cell.paragraphs:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Ensure table is before sectPr, add trailing paragraph
    body = doc.element.body
    sect_pr = body.find(qn('w:sectPr'))
    if sect_pr is not None:
        sect_pr.addprevious(table._tbl)
    doc.add_paragraph()

    _fix_all_table_columns(table)
    
    # Add page break after section
    add_page_break(doc)
    
    print(f"  [OK] Added HPA section")


def add_charge_section_to_docx(
    doc: Document,
    sensors: list | None = None,
    counter: Counter | None = None,
    bridge_number: int = 1,
    precomputed: dict | None = None,
    pilot_finish: str = "",
) -> None:
    """Append section 1.4.6 (charge status table) at end of document body.
    Only processes sensors with type 'gateway' (GW sensors).
    Skips section if no gateway sensors exist.
    """
    from datetime import datetime

    # Filter for gateway sensors (GW sensors)
    gateway_sensors = [s for s in sensors if (s.get("type") or "").lower() in ["gateway", "gateaway", "gw", "snmx"]] if sensors else []
    
    if not gateway_sensors:
        print(f"     ! No gateway sensors found, skipping charge section")
        return


    p = doc.add_paragraph()
    p.style = "Heading 5"
    run = p.add_run(f"{bridge_number}.4.6 Σύνοψη ρεύματος φόρτισης")
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)

    p_cap = doc.add_paragraph()
    p_cap.style = "Caption"
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table_num = counter.next_table() if counter else 54
    p_cap.add_run(f"Πίνακας {table_num}: Τελευταίες μετρήσεις φόρτισης για επαναφορτιζόμενες συσκευές").italic = True

    # Create table with header row
    headers = ["Αισθητήρας", "Site", "Ώρα", "Ρεύμα φόρτισης (mA)", "Status", "Notes"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header row
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        cell = hdr_cells[i]
        cell.text = header
        for para in cell.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                run.font.bold = True
                run.font.size = Pt(10)
        tcPr = cell._element.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), 'B4C6E7')
        tcPr.append(shd)

    # Add row for each gateway sensor
    print(f"     DEBUG charge section: processing {len(gateway_sensors)} gateway sensors")
    charge_results = _batch_check_sensors(gateway_sensors, check_charge, precomputed=precomputed)
    for sensor in gateway_sensors:
        device_id = sensor.get("device_id", "")
        # Try multiple fields for sensor name from bridge API
        sensor_name = (sensor.get("Code_name", "") or 
                       sensor.get("name", "") or 
                       sensor.get("sensor_name", "") or
                       sensor.get("sensor_code", "") or
                       sensor.get("device_id", ""))
        site_id = sensor.get("site_id", "")
        print(f"     DEBUG charge sensor: device_id={device_id}, name={sensor_name}, site_id={site_id}, type={sensor.get('type')}")

        if device_id:
            result = charge_results.get(device_id, {})
            if result.get("ok"):
                charge_value = result.get("metrics", {}).get("Charge", "")
                status = "Καλή"
                _explain = result.get("explain", "")
                _upper_words = re.findall(r'\b[Α-ΩΆΈΉΊΌΎΏΪΫ]{2,}\b', _explain)
                comments = " ".join(_upper_words) if _upper_words else _explain
                latest_time = result.get("latest_time", "")
            else:
                charge_value = ""
                status = "N/A"
                comments = result.get("error", "Σφάλμα ελέγχου")
                latest_time = ""
        else:
            charge_value = ""
            status = "N/A"
            comments = "Δεν υπάρχει device_id"
            latest_time = ""

        # Use pilot finish date if no specific timestamp available
        if not latest_time:
            latest_time = _format_date_el(pilot_finish) or datetime.now().strftime("%d/%m/%Y")
        else:
            latest_time = _format_date_el(latest_time)

        # Fallback to device_id if sensor_name is still empty
        if not sensor_name:
            sensor_name = device_id

        row_cells = table.add_row().cells
        row_cells[0].text = shorten_sensor_name(sensor_name)
        row_cells[1].text = site_id
        row_cells[2].text = latest_time
        # Ensure charge is rounded to 3 decimals
        if charge_value and isinstance(charge_value, (int, float)):
            charge_value = f"{float(charge_value):.3f}"
        row_cells[3].text = str(charge_value) if charge_value else ""
        row_cells[4].text = status
        row_cells[5].text = comments

        for cell in row_cells:
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Ensure table is before sectPr, add trailing paragraph
    body = doc.element.body
    sect_pr = body.find(qn('w:sectPr'))
    if sect_pr is not None:
        sect_pr.addprevious(table._tbl)
    doc.add_paragraph()

    _fix_all_table_columns(table)
    
    # Add page break after section
    add_page_break(doc)
    
    print(f"  [OK] Added charge section")


def add_wls_section_to_docx(
    doc: Document,
    sensors: list | None = None,
    counter: Counter | None = None,
    bridge_number: int = 1,
    precomputed: dict | None = None,
    pilot_finish: str = "",
) -> None:
    """Append section 1.4.7 (water level sensor table) at end of document body.
    Only processes sensors with type 'waterlevel' (WL sensors).
    Skips section if no waterlevel sensors exist.
    """
    from datetime import datetime

    # Filter for waterlevel sensors (WL sensors)
    wls_sensors = [s for s in sensors if (s.get("type") or "").lower() == "waterlevel"] if sensors else []
    
    if not wls_sensors:
        print(f"     ! No waterlevel sensors found, skipping WLS section")
        return


    p = doc.add_paragraph()
    p.style = "Heading 5"
    run = p.add_run(f"{bridge_number}.4.7 Αισθητήρες στάθμης υδάτων")
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)

    p_cap = doc.add_paragraph()
    p_cap.style = "Caption"
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table_num = counter.next_table() if counter else 55
    p_cap.add_run(f"Πίνακας {table_num}: Τελευταίες μετρήσεις στάθμης υδάτων").italic = True

    # Create table with header row
    headers = ["Αισθητήρας", "Site", "Ώρα", "Στάθμη ύδατος", "Τελευταία μέτρηση", "Status", "Σχόλια"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header row
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        cell = hdr_cells[i]
        cell.text = header
        for para in cell.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                run.font.bold = True
                run.font.size = Pt(10)
        tcPr = cell._element.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), 'B4C6E7')
        tcPr.append(shd)

    # Add row for each waterlevel sensor
    print(f"     DEBUG WLS section: processing {len(wls_sensors)} waterlevel sensors")
    wls_results = _batch_check_sensors(wls_sensors, check_wls, precomputed=precomputed)
    for sensor in wls_sensors:
        device_id = sensor.get("device_id", "")
        # Try multiple fields for sensor name from bridge API
        sensor_name = (sensor.get("Code_name", "") or 
                       sensor.get("name", "") or 
                       sensor.get("sensor_name", "") or
                       sensor.get("sensor_code", "") or
                       sensor.get("device_id", ""))
        site_id = sensor.get("site_id", "")
        print(f"     DEBUG WLS sensor: device_id={device_id}, name={sensor_name}, site_id={site_id}, type={sensor.get('type')}")

        if device_id:
            result = wls_results.get(device_id, {})
            if result.get("ok"):
                water_level = result.get("metrics", {}).get("Water Level", "")
                latest = result.get("metrics", {}).get("Latest", "")
                status = "Καλή"
                _explain = result.get("explain", "")
                _upper_words = re.findall(r'\b[Α-ΩΆΈΉΊΌΎΏΪΫ]{2,}\b', _explain)
                comments = " ".join(_upper_words) if _upper_words else _explain
                latest_time = result.get("latest_time", "")
            else:
                water_level = ""
                latest = ""
                status = "N/A"
                comments = result.get("error", "Σφάλμα ελέγχου")
                latest_time = ""
        else:
            water_level = ""
            latest = ""
            status = "N/A"
            comments = "Δεν υπάρχει device_id"
            latest_time = ""

        # Use pilot finish date if no specific timestamp available
        if not latest_time:
            latest_time = _format_date_el(pilot_finish) or datetime.now().strftime("%d/%m/%Y")
        else:
            latest_time = _format_date_el(latest_time)

        # Fallback to device_id if sensor_name is still empty
        if not sensor_name:
            sensor_name = device_id

        row_cells = table.add_row().cells
        row_cells[0].text = shorten_sensor_name(sensor_name)
        row_cells[1].text = site_id
        row_cells[2].text = latest_time
        # Ensure water_level and latest are rounded to 3 decimals
        if water_level and isinstance(water_level, (int, float)):
            water_level = f"{float(water_level):.3f}"
        if latest and isinstance(latest, (int, float)):
            latest = f"{float(latest):.3f}"
        row_cells[3].text = str(water_level) if water_level else ""
        row_cells[4].text = str(latest) if latest else ""
        row_cells[5].text = status
        row_cells[6].text = comments

        for cell in row_cells:
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Ensure table is before sectPr, add trailing paragraph
    body = doc.element.body
    sect_pr = body.find(qn('w:sectPr'))
    if sect_pr is not None:
        sect_pr.addprevious(table._tbl)
    doc.add_paragraph()

    _fix_all_table_columns(table)
    
    # Add page break after section
    add_page_break(doc)
    
    print(f"  [OK] Added WLS section")


def cli_main() -> None:
    ap = argparse.ArgumentParser(description="Generate bridge reports from API + Excel.")
    ap.add_argument("--excel", required=False, default=None, type=Path, help="Path to .xlsx with code_names.")
    ap.add_argument("--template", required=True, type=Path, help="Path to .docx template.")
    ap.add_argument("--out", type=Path, default=Path("output"), help="Output directory.")
    ap.add_argument("--column", default="1",
                    help="Column containing code_name (1-based index or letter). Default: 1")
    ap.add_argument("--sheet", default=None, help="Worksheet name. Default: active sheet.")
    ap.add_argument("--no-header", action="store_true",
                    help="Set if the first row is data (no header row).")
    ap.add_argument("--dump-tags", action="store_true",
                    help="Also write a .tags.json file per bridge for debugging.")
    ap.add_argument("--api-code-name", default=None,
                    help="Fetch site IDs from bridge API instead of Excel (e.g. ΚΜΑ_006_Α07).")
    args = ap.parse_args()

    column: str | int = int(args.column) if args.column.isdigit() else args.column
    process(
        excel=args.excel,
        template=args.template,
        out_dir=args.out,
        column=column,
        sheet=args.sheet,
        skip_header=not args.no_header,
        dump_tags=args.dump_tags,
        api_code_name=args.api_code_name,
    )


# ---------------------------------------------------------------------------
# Multi-bridge report generator
# ---------------------------------------------------------------------------
def read_bridges_txt(txt_path: Path) -> list[str]:
    """Read one code_name per line from a .txt file, skip empty lines and comments."""
    lines = txt_path.read_text(encoding="utf-8-sig").splitlines()
    return [l.strip().rstrip(".") for l in lines if l.strip() and not l.strip().startswith("#")]


def _copy_image_part(target_doc: Document, source_part, r_id: str) -> str:
    """Copy an image part from source to target document and return the new rId."""
    import io
    image_stream = io.BytesIO(source_part._blob)
    new_part = target_doc.part.package.get_or_add_image_part(image_stream)
    new_rid = target_doc.part.relate_to(new_part, docx.opc.constants.RELATIONSHIP_TYPE.IMAGE)
    return new_rid


def _copy_styles_from_source(target: Document, source: Document) -> None:
    """Copy paragraph styles from source to target that don't already exist.

    Ensures appended paragraphs referencing Heading 4/5/etc. styles are still
    recognized as headings in the merged document.
    """
    target_styles = target.styles
    target_styles_el = target_styles._element
    existing_ids = {style.style_id for style in target_styles}
    existing_names = {style.name for style in target_styles}

    for style in source.styles:
        if style.type != 1:  # only paragraph styles
            continue
        if style.style_id in existing_ids or style.name in existing_names:
            continue
        # Copy the style element to the target styles part
        new_style_el = copy.deepcopy(style.element)
        target_styles_el.append(new_style_el)
        existing_ids.add(style.style_id)
        existing_names.add(style.name)


def append_docx_to(target: Document, source: Document) -> None:
    """Append all body elements from source into target document.

    Copies embedded image parts and updates relationship IDs so images are not broken.
    """
    # Build map of source part rIds -> target rIds for all relationship types
    source_rid_map: dict[str, str] = {}
    for rel in source.part.rels.values():
        if rel.reltype == docx.opc.constants.RELATIONSHIP_TYPE.IMAGE:
            source_part = source.part.related_parts[rel.rId]
            target_rid = _copy_image_part(target, source_part, rel.rId)
            source_rid_map[rel.rId] = target_rid
        elif rel.reltype == docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK:
            # Copy external hyperlink relationship
            target_rid = target.part.relate_to(
                rel.target_ref, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True
            )
            source_rid_map[rel.rId] = target_rid

    # Copy body elements, updating relationship references
    for element in source.element.body:
        if element.tag.endswith('}sectPr'):
            continue
        new_element = copy.deepcopy(element)
        # Update all blip embed references to point to new image parts
        NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
        for blip in new_element.findall(f".//{{{NS_A}}}blip"):
            embed = blip.get(qn("r:embed"))
            if embed and embed in source_rid_map:
                blip.set(qn("r:embed"), source_rid_map[embed])
        # Update all hyperlink references
        for link in new_element.findall(f".//{qn('w:hlink')}"):
            rel = link.get(qn("r:id"))
            if rel and rel in source_rid_map:
                link.set(qn("r:id"), source_rid_map[rel])
        target.element.body.append(new_element)


def add_page_break(doc: Document) -> None:
    p = OxmlElement('w:p')
    r = OxmlElement('w:r')
    br = OxmlElement('w:br')
    br.set(qn('w:type'), 'page')
    r.append(br)
    p.append(r)
    body = doc.element.body
    sect_pr = body.find(qn('w:sectPr'))
    if sect_pr is not None:
        sect_pr.addprevious(p)
    else:
        body.append(p)


def generate_report(
    bridges_txt: Path,
    bridge_template: Path,
    main_report: Path | None,
    out_path: Path | None,
    region_name: str = "",
    start_date: str = "",
    end_date: str = "",
) -> None:
    """Read code_names from bridges_txt, fetch API data, render per-bridge sections
    from bridge_template, and insert them into main_report under the section
    'Αναλυτική παρουσίαση αποτελεσμάτων ανά γέφυρα'.

    Args:
        bridges_txt: path to .txt file with one code_name per line
        bridge_template: path to per-bridge .docx template with {placeholders}
        main_report: path to main .docx report
        out_path: output .docx path
        region_name: replaces {REGION_NAME} in all documents
        start_date: replaces {START_DATE}
        end_date: replaces {END_DATE}
    """
    code_names = read_bridges_txt(bridges_txt)
    if not code_names:
        print("No code_names found in bridges.txt.")
        return
    print(f"Found {len(code_names)} bridges: {code_names}")

    # Global tags that apply everywhere
    global_tags = {
        "REGION_NAME": region_name,
        "START_DATE": start_date,
        "END_DATE": end_date,
    }

    out_dir = out_path.parent if out_path else bridges_txt.parent / "output"
    out_dir.mkdir(parents=True, exist_ok=True)

    # Step 1: render and save one .docx per bridge using generate_bridge_docx
    # Process bridges in parallel using ThreadPoolExecutor
    from concurrent.futures import ThreadPoolExecutor, as_completed
    bridge_paths: list[Path] = []
    results: dict = {}
    with ThreadPoolExecutor(max_workers=4) as ex:
        futures = {}
        for idx, code_name in enumerate(code_names, start=1):
            f = ex.submit(
                generate_bridge_docx,
                code_name=code_name,
                template=bridge_template,
                out_dir=out_dir,
                extra_tags=global_tags,
                bridge_number=idx,
            )
            futures[f] = idx
        for f in as_completed(futures):
            idx = futures[f]
            try:
                bp = f.result()
                if bp:
                    results[idx] = bp
            except Exception as exc:
                print(f"  ! Bridge [{idx}] failed: {exc}")
    # Preserve original order
    bridge_paths = [results[i] for i in sorted(results) if i in results]

    print(f"\nPer-bridge files: {[str(p) for p in bridge_paths]}")

    if main_report is None or out_path is None:
        return

    # Step 2: open main report, apply global tags, append all bridge sections
    main_doc = Document(str(main_report))
    for paragraph in _iter_paragraphs(main_doc):
        _replace_in_paragraph(paragraph, global_tags)

    for bridge_path in bridge_paths:
        bridge_doc = Document(str(bridge_path))
        add_page_break(main_doc)
        _strip_trailing_logo(bridge_doc)
        append_docx_to(main_doc, bridge_doc)

    add_ellada2_logo(main_doc)
    main_doc.save(str(out_path))
    try:
        PILOT_REPORTS_DIR.mkdir(parents=True, exist_ok=True)
        shutil.copy2(str(out_path), str(PILOT_REPORTS_DIR / out_path.name))
    except Exception as e:
        print(f"     ! Failed to copy combined report to pilot reports dir: {e}")
    print(f"Report saved:     {out_path}")


def main() -> None:
    ap = argparse.ArgumentParser(description="Generate merged bridge report from bridges.txt.")
    ap.add_argument("--bridges", required=True, type=Path, help="Path to bridges.txt (one code_name per line).")
    ap.add_argument("--template", required=True, type=Path, help="Path to per-bridge .docx template.")
    ap.add_argument("--report", required=False, default=None, type=Path, help="Path to main .docx report. If omitted, only per-bridge files are generated.")
    ap.add_argument("--out", type=Path, default=Path("output/report_final.docx"), help="Output .docx path.")
    ap.add_argument("--region", default="", help="Region name ({REGION_NAME}).")
    ap.add_argument("--start-date", default="", help="Start date ({START_DATE}).")
    ap.add_argument("--end-date", default="", help="End date ({END_DATE}).")
    args = ap.parse_args()

    generate_report(
        bridges_txt=args.bridges,
        bridge_template=args.template,
        main_report=args.report,
        out_path=args.out if args.report else None,
        region_name=args.region,
        start_date=args.start_date,
        end_date=args.end_date,
    )


if __name__ == "__main__":
    import sys
    if "--bridges" in sys.argv:
        main()
    else:
        cli_main()
